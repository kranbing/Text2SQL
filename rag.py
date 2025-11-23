"""
RAG 检索模块
功能：构建向量索引或启用简易检索，返回相关文档。
方法：
- RAGIndex.build(docs): 构建索引
- RAGRetriever.set_index(index): 设置索引
- RAGRetriever.query(question, top_k): 返回相关文档列表
"""
class RAGIndex:
    def __init__(self):
        self.docs = []
        self.vs = None
        try:
            from langchain_openai import OpenAIEmbeddings
            from langchain_community.vectorstores import FAISS
            self._emb_cls = OpenAIEmbeddings
            self._vs_cls = FAISS
        except Exception:
            self._emb_cls = None
            self._vs_cls = None
    def _extract_text_from_file(self, path: str) -> str | None:
        try:
            ext = (path.split('.')[-1] or '').lower()
            if ext in ("txt", "md"):
                with open(path, "r", encoding="utf-8") as f:
                    return f.read()
            if ext == "csv":
                import csv
                out = []
                with open(path, "r", encoding="utf-8", newline="") as f:
                    r = csv.reader(f)
                    for row in r:
                        out.append("\t".join(str(x) for x in row))
                return "\n".join(out)
            if ext in ("xlsx",):
                try:
                    import openpyxl
                    wb = openpyxl.load_workbook(path, read_only=True)
                    out = []
                    for ws in wb.worksheets:
                        for row in ws.iter_rows(values_only=True):
                            out.append("\t".join(str(x) if x is not None else "" for x in row))
                    return "\n".join(out)
                except Exception:
                    return None
            if ext in ("xls",):
                try:
                    import xlrd
                    book = xlrd.open_workbook(path)
                    out = []
                    for si in range(book.nsheets):
                        sh = book.sheet_by_index(si)
                        for ri in range(sh.nrows):
                            row = sh.row_values(ri)
                            out.append("\t".join(str(x) for x in row))
                    return "\n".join(out)
                except Exception:
                    return None
            if ext in ("docx",):
                try:
                    from docx import Document
                    doc = Document(path)
                    return "\n".join(p.text for p in doc.paragraphs)
                except Exception:
                    return None

        except Exception:
            return None
        return None
    def load_docs_from_dir(self, dir_path: str, glob_patterns: str | None = None) -> list[str]:
        import os
        import glob
        if not dir_path or not os.path.isdir(dir_path):
            return []
        pats = []
        if glob_patterns:
            for p in str(glob_patterns).split(","):
                p = p.strip()
                if p:
                    pats.append(p)
        if not pats:
            pats = ["*.txt", "*.md", "*.csv", "*.xlsx", "*.xls", "*.docx"]
        texts = []
        for gp in pats:
            for p in glob.glob(os.path.join(dir_path, gp)):
                t = self._extract_text_from_file(p)
                if t and t.strip():
                    texts.append(t)
        return texts
    def build(self, docs: list[str]):
        self.docs = docs or []
        if self._emb_cls and self._vs_cls and self.docs:
            try:
                import os
                base_url = os.getenv("SILICONFLOW_BASE_URL", None)
                api_key = os.getenv("SILICONFLOW_API_KEY", None)
                kwargs = {}
                if base_url:
                    kwargs["base_url"] = base_url
                if api_key:
                    kwargs["api_key"] = api_key
                emb = self._emb_cls(**kwargs)
                self.vs = self._vs_cls.from_texts(self.docs, emb)
            except Exception:
                self.vs = None

class RAGRetriever:
    def __init__(self):
        self.index = None
    def set_index(self, index: RAGIndex):
        self.index = index
    def query(self, question: str, top_k: int = 5) -> list[str]:
        if not self.index:
            return []
        if self.index.vs is not None:
            try:
                docs = self.index.vs.similarity_search(question, k=top_k)
                return [d.page_content for d in docs]
            except Exception:
                pass
        q = question.lower()
        scored = []
        for d in self.index.docs:
            s = 0
            for tok in q.split():
                if tok and tok in d.lower():
                    s += 1
            scored.append((s, d))
        scored.sort(key=lambda x: x[0], reverse=True)
        return [d for s, d in scored[:top_k] if s > 0]